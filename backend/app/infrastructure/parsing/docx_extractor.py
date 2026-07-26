"""
DOCX text and structure extraction via python-docx.

Known limitation (documented, not hidden): python-docx has no
reliable page-count API without actually rendering the document (that
would require LibreOffice/docx2pdf in the container, which we
deliberately avoid to keep this dependency-light and fast). We report
page_count=1 as a "not confidently determined" default rather than
guessing from paragraph count, which would be misleading precision.
"""

from __future__ import annotations

import io

import docx

from app.domain.entities.resume import FormattingFlags
from app.infrastructure.parsing.raw_document import RawDocument

_STANDARD_FONT_KEYWORDS = (
    "arial", "helvetica", "times", "calibri", "cambria", "georgia",
    "verdana", "garamond", "tahoma", "trebuchet", "segoe",
)


def extract_docx(file_bytes: bytes) -> RawDocument:
    document = docx.Document(io.BytesIO(file_bytes))

    paragraph_texts = [p.text for p in document.paragraphs]

    table_texts = []
    for table in document.tables:
        for row in table.rows:
            table_texts.append(" | ".join(cell.text for cell in row.cells))

    full_text = "\n".join(paragraph_texts + table_texts)

    uses_tables = len(document.tables) > 0
    uses_text_boxes = _has_text_boxes(document)
    non_standard_fonts = _has_non_standard_fonts(document)
    has_images = len(document.inline_shapes) > 0

    flags = FormattingFlags(
        uses_tables=uses_tables,
        uses_text_boxes=uses_text_boxes,
        uses_columns=False,
        non_standard_fonts=non_standard_fonts,
        has_images_with_text=has_images,
        page_count=1,  # see module docstring
    )
    return RawDocument(text=full_text, formatting_flags=flags)


def _has_text_boxes(document: docx.Document) -> bool:
    for element in document.element.body.iter():
        tag = element.tag
        if isinstance(tag, str) and tag.endswith("}txbxContent"):
            return True
    return False


def _has_non_standard_fonts(document: docx.Document) -> bool:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            font_name = (run.font.name or "").lower()
            if font_name and not any(
                kw in font_name for kw in _STANDARD_FONT_KEYWORDS
            ):
                return True
    return False
