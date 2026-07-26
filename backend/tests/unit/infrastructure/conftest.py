from __future__ import annotations

import io

import docx
import pytest
from fpdf import FPDF

_SAMPLE_RESUME_LINES = [
    "Jane Smith",
    "jane.smith@example.com | 555-123-4567 | linkedin.com/in/janesmith",
    "",
    "Summary",
    "Backend engineer with a focus on distributed systems.",
    "",
    "Experience",
    "Backend Engineer, Acme Corp",
    "Jan 2021 - Jun 2023",
    "- Built a scalable REST API serving 10k requests/sec using FastAPI",
    "- Reduced database query latency by 40 percent through indexing",
    "",
    "Software Engineer, Beta Inc",
    "Aug 2023 - Present",
    "- Migrated monolith to microservices architecture",
    "",
    "Education",
    "Bachelor of Science in Computer Science",
    "State University",
    "2020",
    "",
    "Skills",
    "Python, JavaScript, Docker, Kubernetes",
    "",
    "Projects",
    "Open Source CLI Tool",
    "A CLI tool for automating deployments",
    "Technologies: Python, Docker",
    "- Adopted by 500 users",
    "",
    "Certifications",
    "AWS Certified Developer - Amazon - 2022",
]

_SAMPLE_JD_LINES = [
    "Backend Engineer",
    "",
    "Responsibilities",
    "- Design and build scalable REST APIs",
    "- Collaborate with product and design teams",
    "",
    "Required",
    "- Python",
    "- JavaScript",
    "- 3+ years of experience",
    "- Bachelor's degree in Computer Science",
    "",
    "Preferred",
    "- Kubernetes",
    "- AWS",
]


@pytest.fixture
def sample_resume_text() -> str:
    return "\n".join(_SAMPLE_RESUME_LINES)


@pytest.fixture
def sample_jd_text() -> str:
    return "\n".join(_SAMPLE_JD_LINES)


@pytest.fixture
def sample_resume_docx_bytes() -> bytes:
    document = docx.Document()
    for line in _SAMPLE_RESUME_LINES:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_resume_docx_with_table_bytes() -> bytes:
    document = docx.Document()
    for line in _SAMPLE_RESUME_LINES[:5]:
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Years"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "5"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_resume_pdf_bytes() -> bytes:
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in _SAMPLE_RESUME_LINES:
        pdf.multi_cell(
            0,
            6,
            text=line if line.strip() else " ",
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )
    return bytes(pdf.output())
